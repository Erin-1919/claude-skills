"""Checks for the poster build eval. Run in the eval workspace.

python check_poster.py <check> [W H]
checks: size W H | pdf | figure | table | headline | minfont | nooverflow
Exit 0 = pass.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ws = Path.cwd()
pptx = sorted(p for p in ws.rglob("*.pptx") if not p.name.startswith("~$"))
if not pptx:
    sys.exit("no .pptx in workspace")
prs = Presentation(pptx[0])
slide = prs.slides[0]
shapes = list(slide.shapes)
check = sys.argv[1]


def inch(v):
    return Emu(v).inches


def runs():
    for s in shapes:
        frames = []
        if s.has_text_frame:
            frames.append(s.text_frame)
        if getattr(s, "has_table", False) and s.has_table:
            frames += [c.text_frame for r in s.table.rows for c in r.cells]
        for tf in frames:
            for p in tf.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        yield r


if check == "size":
    w, h = float(sys.argv[2]), float(sys.argv[3])
    got = (inch(prs.slide_width), inch(prs.slide_height))
    ok = abs(got[0] - w) < 0.1 and abs(got[1] - h) < 0.1
    sys.exit(0 if ok else f"slide is {got[0]:.2f} x {got[1]:.2f} in")
elif check == "pdf":
    sys.exit(0 if list(ws.rglob("*.pdf")) else "no PDF")
elif check == "figure":
    pics = [s for s in shapes if s.shape_type == 13]
    big = [s for s in pics if inch(s.width) >= 10]
    sys.exit(0 if big else f"{len(pics)} pictures, none >= 10 in wide")
elif check == "table":
    sys.exit(0 if any(getattr(s, "has_table", False) and s.has_table for s in shapes)
             else "no native table")
elif check == "headline":
    text = " ".join(r.text for r in runs()).lower()
    sys.exit(0 if "topology carries the reasoning" in text else "headline missing")
elif check == "minfont":
    sizes = [r.font.size.pt for r in runs() if r.font.size is not None]
    if not sizes:
        sys.exit("no explicit font sizes")
    small = [s for s in sizes if s < 20]
    sys.exit(0 if not small else f"{len(small)} runs under 20 pt (min {min(sizes)})")
elif check == "nooverflow":
    W, H = prs.slide_width, prs.slide_height
    out = [s.name for s in shapes
           if s.left < 0 or s.top < 0 or s.left + s.width > W + 1 or s.top + s.height > H + 1]
    sys.exit(0 if not out else f"off-slide: {out}")
sys.exit("unknown check " + check)
