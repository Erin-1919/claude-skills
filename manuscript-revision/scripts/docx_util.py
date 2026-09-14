# -*- coding: utf-8 -*-
"""Shared python-docx helpers for editing a manuscript in place.

Two traps this works around, both of which cost real time:

1. `Paragraph.runs` skips runs inside `w:hyperlink`, so a paragraph containing a link
   reads short and an edit silently misses text. `all_runs` walks `w:r` directly.
2. A phrase is almost never in one run. Word splits runs at spell-check boundaries, at
   formatting changes, and after any edit. `splice` replaces a span that straddles runs
   while keeping the formatting of the run where the span starts.

Import from a script in this folder:

    import sys, os
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from docx_util import all_runs, ptext, splice, iter_paragraphs, find_para
"""
import re

import docx
from docx.text.run import Run

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'


def all_runs(paragraph):
    """Every run, including runs inside hyperlinks."""
    return [Run(r, paragraph) for r in paragraph._p.iter(W + 'r')]


def ptext(paragraph):
    return ''.join(r.text for r in all_runs(paragraph))


def iter_paragraphs(doc, tables=True):
    """Body paragraphs, then table-cell paragraphs."""
    for p in doc.paragraphs:
        yield p
    if tables:
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p


def iter_document_order(doc):
    """Paragraphs in true document order, descending into tables where they sit.

    `doc.paragraphs` then `doc.tables` puts every table after all prose, which breaks any
    check that depends on the order things appear, such as first-appearance citation
    ordering. This walks the body's children instead.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def walk(parent, element):
        for child in element:
            if child.tag == W + 'p':
                yield Paragraph(child, parent)
            elif child.tag == W + 'tbl':
                table = Table(child, parent)
                for row in table.rows:
                    for cell in row.cells:
                        for p in walk(cell, cell._tc):
                            yield p

    for p in walk(doc, doc.element.body):
        yield p


def doc_text(path, tables=True):
    """All visible text of a file as one string, for existence checks."""
    doc = docx.Document(path)
    return ' '.join(ptext(p) for p in iter_paragraphs(doc, tables))


def find_para(doc, needle, tables=False):
    """The one paragraph containing needle. Raises unless exactly one matches."""
    hits = [p for p in iter_paragraphs(doc, tables) if needle in ptext(p)]
    if len(hits) != 1:
        raise ValueError('%d paragraphs contain %r' % (len(hits), needle[:60]))
    return hits[0]


def splice(paragraph, old, new):
    """Replace the first occurrence of old, even when it straddles runs.

    The replacement takes the formatting of the run where old begins.
    """
    runs = all_runs(paragraph)
    texts = [r.text for r in runs]
    full = ''.join(texts)
    if old not in full:
        raise ValueError('not found: %r' % old[:60])
    i = full.index(old)
    j = i + len(old)
    pos = 0
    for run, text in zip(runs, texts):
        start, end = pos, pos + len(text)
        pos = end
        if end <= i or start >= j:
            continue
        a = max(start, i) - start
        b = min(end, j) - start
        run.text = text[:a] + (new if start <= i < end else '') + text[b:]
    return True


def replace_all(paragraph, old, new):
    n = 0
    while old in ptext(paragraph):
        splice(paragraph, old, new)
        n += 1
    return n


def set_text(paragraph, text):
    """Replace a paragraph's whole text, keeping the first run's formatting.

    Only safe when the paragraph's runs share one format. Check with uniform().
    """
    runs = all_runs(paragraph)
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def uniform(paragraph):
    """True when every run shares bold, italic, superscript, font and size."""
    sig = {(r.bold, r.italic, r.font.superscript, r.font.name, r.font.size)
           for r in all_runs(paragraph)}
    return len(sig) <= 1


def has_math(paragraph):
    """True when the paragraph holds an OMML equation, whose text `ptext` cannot see."""
    return paragraph._p.find('.//' + M + 'oMath') is not None


def style_run(run, font=None, size=None, color=None, bold=None):
    """Apply the letter's own formatting to a new run."""
    from docx.shared import Pt, RGBColor
    if font:
        run.font.name = font
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(int(color[0:2], 16), int(color[2:4], 16),
                                      int(color[4:6], 16))
    if bold is not None:
        run.bold = bold
    return run


def clone_paragraph_after(paragraph, text=None):
    """A new empty paragraph with the same properties, inserted after this one."""
    import copy
    from docx.text.paragraph import Paragraph
    el = copy.deepcopy(paragraph._p)
    for r in el.findall(W + 'r'):
        el.remove(r)
    paragraph._p.addnext(el)
    new = Paragraph(el, paragraph._parent)
    if text is not None:
        new.add_run(text)
    return new


def normalize(s):
    """Collapse whitespace and curly apostrophes, for comparing quoted text."""
    return re.sub(r'\s+', ' ', s.replace('’', "'").replace('‘', "'")).strip()
