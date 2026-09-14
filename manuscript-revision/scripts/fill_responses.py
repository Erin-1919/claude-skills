# -*- coding: utf-8 -*-
"""Write answers from answers.py into the letter's placeholders.

Run after adding an entry to answers.py. Only placeholders are touched, so a response
already written, including one edited by hand in Word, is never overwritten.

    python fill_responses.py            # dry run, says what it would fill
    python fill_responses.py --write

Matching: an entry's key must appear in the placeholder text. Placeholders read
"[TO BE COMPLETED - R1.1]", so a key of "R1.1" matches. Where several placeholders would
match one key, the first unfilled one is used, so keys should be distinctive.
"""
import os
import shutil
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import revision_config
from docx_util import all_runs, clone_paragraph_after, ptext
from letter_io import TODO, counts, load_answers, styled


def fill(doc, key, body, cfg):
    target = next((p for p in doc.paragraphs
                   if TODO in ptext(p) and key in ptext(p)), None)
    if target is None:
        return False
    for r in all_runs(target):
        r._r.getparent().remove(r._r)
    styled(target, 'Response: ', cfg, bold=True)
    styled(target, body[0], cfg)
    anchor = target
    for text in body[1:]:
        anchor = clone_paragraph_after(anchor)
        styled(anchor, text, cfg)
    return True


if __name__ == '__main__':
    cfg = revision_config.load()
    answers = load_answers(cfg['answers'])
    doc = docx.Document(cfg['letter'])
    filled = 0
    for key, body in answers.items():
        ok = fill(doc, key, body, cfg)
        filled += ok
        print('%s %s' % ('filled   ' if ok else 'not found', key))
    left = [ptext(p) for p in doc.paragraphs if TODO in ptext(p)]
    print('\nfilled %d | placeholders remaining: %d' % (filled, len(left)))
    for t in left:
        print('   %s' % t.strip()[:100])
    if '--write' in sys.argv:
        shutil.copyfile(cfg['letter'], cfg['letter'] + '.bak')
        doc.save(cfg['letter'])
        print('\nwritten | placeholders, written = %s' % (counts(cfg['letter']),))
    else:
        print('\ndry run - pass --write to apply')
