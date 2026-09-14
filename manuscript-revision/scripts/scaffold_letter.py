# -*- coding: utf-8 -*-
"""Build the response letter from the reviewer comments, with placeholders.

Do this once, at the start, so every comment is in the letter verbatim before any of them
is answered. Nothing is ever retyped from the decision email afterwards.

Input is a plain markdown file, comments.md:

    # Editor
    <optional editor text>

    # Reviewer 1
    <the reviewer's overall assessment, one or more paragraphs>

    ## R1.1
    <the comment, verbatim, one or more paragraphs>

    ## R1.2
    ...

Output is the letter at the path in revision.json, with a heading per reviewer, each
comment in the body style, and a placeholder response after each. The comment ids become
the keys used everywhere else, so keep them short and stable.

    python scaffold_letter.py comments.md
    python scaffold_letter.py comments.md --write
"""
import io
import os
import re
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import revision_config
from letter_io import styled

PLACEHOLDER = '[TO BE COMPLETED - %s]'


def parse(path):
    """Return [(heading, overall_paragraphs, [(id, comment_paragraphs), ...]), ...]."""
    sections, current, comment = [], None, None
    for raw in io.open(path, encoding='utf-8'):
        line = raw.rstrip('\n')
        if line.startswith('# '):
            current = (line[2:].strip(), [], [])
            sections.append(current)
            comment = None
        elif line.startswith('## '):
            if current is None:
                raise SystemExit('comment %r appears before any "# Reviewer" heading'
                                 % line)
            comment = (line[3:].strip(), [])
            current[2].append(comment)
        elif line.strip():
            if comment is not None:
                comment[1].append(line.strip())
            elif current is not None:
                current[1].append(line.strip())
    if not sections:
        raise SystemExit('no "# Editor" or "# Reviewer N" headings found in %s' % path)
    return sections


def build(sections, cfg, out_path):
    doc = docx.Document()
    for name, overall, comments in sections:
        head = doc.add_paragraph()
        styled(head, 'Response to %s' % (name if name.lower() == 'editor'
                                         else '%s Comments' % name), cfg, bold=True)
        for text in overall:
            styled(doc.add_paragraph(), text, cfg)
        label = ('editor page, written last' if name.lower() == 'editor'
                 else 'overall assessment for %s, written last' % name)
        p = doc.add_paragraph()
        styled(p, 'Response: ', cfg, bold=True)
        styled(p, PLACEHOLDER % label, cfg)
        for cid, body in comments:
            first = doc.add_paragraph()
            styled(first, 'Comment %s: ' % cid, cfg, bold=True)
            styled(first, body[0] if body else '', cfg)
            for text in body[1:]:
                styled(doc.add_paragraph(), text, cfg)
            p = doc.add_paragraph()
            styled(p, 'Response: ', cfg, bold=True)
            styled(p, PLACEHOLDER % cid, cfg)
    doc.save(out_path)


def answers_stub(path, sections):
    if os.path.exists(path):
        print('answers file already exists, left alone: %s' % path)
        return
    ids = [cid for _, _, comments in sections for cid, _ in comments]
    io.open(path, 'w', encoding='utf-8').write(
        u'# -*- coding: utf-8 -*-\n'
        u'"""One entry per response, keyed by a distinctive fragment of the placeholder.\n\n'
        u'Filled in as each plan item lands, never all at the end. sync_letter.py rewrites\n'
        u'entries here from the letter, so hand edits made in Word are never lost.\n\n'
        u'Comment ids in this round: %s\n"""\n\nANSWERS = {\n}\n' % ', '.join(ids))
    print('wrote answers stub: %s' % path)


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        raise SystemExit(__doc__)
    cfg = revision_config.load()
    sections = parse(args[0])
    total = sum(len(c) for _, _, c in sections)
    for name, overall, comments in sections:
        print('  %-14s overall paragraphs: %-2d comments: %d'
              % (name, len(overall), len(comments)))
    print('\n%d sections, %d comments, %d placeholders'
          % (len(sections), total, total + len(sections)))
    if '--write' not in sys.argv:
        print('\ndry run - pass --write to create the letter')
        sys.exit(0)
    if os.path.exists(cfg['letter']):
        raise SystemExit('letter already exists, refusing to overwrite: %s' % cfg['letter'])
    build(sections, cfg, cfg['letter'])
    print('\nwrote %s' % cfg['letter'])
    answers_stub(cfg['answers'], sections)
