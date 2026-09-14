# -*- coding: utf-8 -*-
"""Shared reading and writing for the response letter and its answers file.

The letter is the artifact the journal sees. `answers.py` is a plain data file holding
one entry per response, keyed by a distinctive fragment of the placeholder or of the first
sentence. The two must never drift apart, which is what sync_letter.py enforces.

Letter structure this assumes, produced by scaffold_letter.py:

    Response to Editor                  <- heading
    Response: ...                       <- response block, one or more paragraphs
    Response to Reviewer 1 Comments     <- heading
    <reviewer's overall assessment>
    Response: ...                       <- per-reviewer summary
    Comment R1.1: <verbatim>            <- one or more paragraphs
    Response: ...                       <- response block
"""
import io
import os
import re
import sys
import textwrap

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_util import all_runs, ptext, style_run

BS = chr(92)
HEAD = 'Response to '
COMMENT = re.compile(r'^Comment ([A-Za-z0-9_.\-]+)\s*:')
TODO = 'TO BE COMPLETED'


def paragraphs(path):
    doc = docx.Document(path)
    return doc, [ptext(p) for p in doc.paragraphs]


def blocks(path):
    """Every written response, keyed by the comment id or by the heading it answers.

    The letter is split into segments, one per heading or comment. Within a segment the
    response is the paragraph beginning "Response: " and the paragraphs that follow it.
    A segment whose response is still a placeholder yields nothing.

    The editor page may carry no "Response: " prefix, following the convention of some
    letters. Only when a heading segment contains no "Response: " paragraph at all is its
    first block of prose taken as the response, which keeps a reviewer's own assessment
    from being mistaken for one.
    """
    _, texts = paragraphs(path)

    marks = []
    for i, t in enumerate(texts):
        m = COMMENT.match(t)
        if m:
            marks.append((i, 'comment', m.group(1)))
        elif t.startswith(HEAD):
            marks.append((i, 'head', t[len(HEAD):].replace(' Comments', '').strip()))

    out = {}
    for k, (start, kind, tag) in enumerate(marks):
        stop = marks[k + 1][0] if k + 1 < len(marks) else len(texts)
        body_start = None
        for i in range(start + 1, stop):
            if texts[i].startswith('Response: '):
                body_start = None if TODO in texts[i] else i
                break
        else:
            if kind == 'head':
                body_start = next((i for i in range(start + 1, stop)
                                   if texts[i].strip()), None)
        if body_start is None:
            continue
        first = texts[body_start]
        body = [first[len('Response: '):] if first.startswith('Response: ') else first]
        j = body_start + 1
        while j < stop and texts[j].strip() and not texts[j].startswith('Response: '):
            body.append(texts[j])
            j += 1
        out[tag] = [p.strip() for p in body]
    return out


def escape(text):
    return ''.join(c if ord(c) < 128 else BS + 'u%04x' % ord(c) for c in text)


def literal(body, indent=8):
    """Render a list of paragraphs as python source, wrapped and escaped."""
    out = []
    for text in body:
        chunks = textwrap.wrap(text, 84, break_long_words=False, break_on_hyphens=False)
        lines = []
        for i, chunk in enumerate(chunks):
            s = chunk if i == len(chunks) - 1 else chunk + ' '
            s = s.replace(BS, BS * 2).replace("'", BS + "'")
            lines.append(' ' * indent + "'%s'" % escape(s))
        out.append('\n'.join(lines) + ',')
    return '\n'.join(out)


def load_answers(path):
    import importlib.util
    spec = importlib.util.spec_from_file_location('answers', path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.ANSWERS


def marker_for(src, key):
    """Find a key's opening line, whether it was stored raw or escaped."""
    for form in ("    '%s': [\n" % key, "    '%s': [\n" % escape(key)):
        if form in src:
            return form
    raise KeyError(key)


def rewrite_answer(path, key, body):
    """Replace one entry's paragraphs in answers.py, leaving the rest of the file alone."""
    src = io.open(path, encoding='utf-8').read()
    marker = marker_for(src, key)
    a = src.index(marker) + len(marker)
    b = src.index('\n    ],\n', a)
    io.open(path, 'w', encoding='utf-8').write(src[:a] + literal(body) + src[b:])


def append_answer(path, key, body, comment=None):
    src = io.open(path, encoding='utf-8').read()
    entry = ''
    if comment:
        entry += u'    # %s\n' % comment
    entry += u"    '%s': [\n%s\n    ],\n" % (escape(key), literal(body))
    i = src.rindex('}')
    io.open(path, 'w', encoding='utf-8').write(src[:i] + entry + src[i:])


def styled(paragraph, text, cfg, bold=False):
    s = cfg.get('letter_style', {})
    return style_run(paragraph.add_run(text), font=s.get('font'),
                     size=s.get('size_pt'), color=s.get('color'), bold=bold)


def counts(path):
    _, texts = paragraphs(path)
    return (sum(1 for t in texts if TODO in t),
            sum(1 for t in texts if t.startswith('Response:') and TODO not in t))
